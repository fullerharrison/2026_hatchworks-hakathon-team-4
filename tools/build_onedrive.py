"""Build the shared UC4 OneDrive folder from the hackathon workspace.

Usage: uv run --no-project --with python-docx --with openpyxl python build_onedrive.py <project_root> <out_dir> [--force]
Refuses to touch an existing <out_dir> unless --force: once shared, teammates edit it and a rebuild would delete their work.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import unquote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(sys.argv[1]).resolve()
OUT = Path(sys.argv[2]).resolve()
FORCE = "--force" in sys.argv[3:]
GS = ROOT / "get_started"
EDA = ROOT / "analysis" / "uc4_eda"
V2_ZIP = "RE__Hatchworks_Hackathon_-_4th_Use_Case.zip"
V1_ZIP = "UC4 - R&D Data Source Unification-20260928T220200Z-1-001.zip"
BRIEFS = "02_Source_Material/Briefs_and_Rules"

# Where each workspace file referenced by the Markdown docs lives in the new tree.
LINK_MAP = {
    "2026_Use_Case_Briefs.pdf": f"{BRIEFS}/2026_Use_Case_Briefs.pdf",
    "2026_Participant_Handbook.pdf": f"{BRIEFS}/2026_Participant_Handbook.pdf",
    "2026_Kickoff_Deck.pdf": f"{BRIEFS}/2026_Kickoff_Deck.pdf",
    V2_ZIP: f"02_Source_Material/Data_v2_CURRENT/{V2_ZIP}",
    V1_ZIP: f"02_Source_Material/Data_v1_SUPERSEDED/{V1_ZIP}",
    "report.html": "04_Analysis/UC4_Data_Profile.html",
    "report_v1.html": "04_Analysis/v1_superseded/report_v1.html",
    "rules.py": "04_Analysis/eda_code/rules.py",
    "uc4_eda.py": "04_Analysis/eda_code/uc4_eda.py",
    "SME_ANSWERS.md": "03_SME/SME_Questions_and_Answers.docx",
    "ROLES.md": "01_Team/Roles_and_Contracts.docx",
    "PROMPT_LOG.md": "06_AI_Prompt_Log/Prompt_Log.xlsx",
    "RECOMMENDATION.md": "04_Analysis/Team_Recommendation.docx",
    "VERIFICATION.md": "04_Analysis/Use_Case_Verification.docx",
    "DATA_ARCHITECTURE.md": "04_Analysis/UC4_Data_Architecture.md",
    "README.md": "04_Analysis/UC4_Guide.docx",
    "v1": "04_Analysis/v1_superseded",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def rewrite_links(md: str, doc_rel: str) -> str:
    """Point local Markdown links at the file's new location, relative to the output doc."""
    doc_dir = os.path.dirname(doc_rel)

    def target(t: str) -> str | None:
        """New relative target, or None for a local file that is not shipped."""
        if re.match(r"^(https?:|mailto:|#)", t):
            return t
        path = unquote(t.split("#")[0]).rstrip("/")
        new = LINK_MAP.get(Path(path).name)
        return os.path.relpath(new, doc_dir or ".").replace("\\", "/") if new else None

    def inline(m: re.Match) -> str:
        t = target(m.group(2))
        return f"[{m.group(1)}]({t})" if t else m.group(1)

    md = re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)", inline, md)
    return re.sub(r"^(\[[^\]]+\]:\s*)(\S+)", lambda m: m.group(1) + (target(m.group(2)) or m.group(2)), md, flags=re.M)


def pandoc(md: str, rel: str, fmt_from: str = "gfm") -> None:
    dest = OUT / rel
    dest.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as f:
        f.write(md)
    try:
        subprocess.run(["pandoc", "-f", fmt_from, "-o", str(dest), f.name], check=True)
    finally:
        os.unlink(f.name)


def convert_doc(src: Path, rel: str) -> None:
    pandoc(rewrite_links(src.read_text(encoding="utf-8"), rel), rel)


def copy(src: Path, rel: str) -> None:
    dest = OUT / rel
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)


def note(rel_dir: str, text: str) -> None:
    """Browser uploads skip empty folders, so every placeholder folder gets a short note."""
    d = OUT / rel_dir
    d.mkdir(parents=True, exist_ok=True)
    (d / "_README.txt").write_text(text + "\n", encoding="utf-8")


def sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def dropdown(ws, col: str, values: list[str]) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}500")


# ---------------------------------------------------------------- source material
def build_source() -> None:
    for pdf in ("2026_Use_Case_Briefs.pdf", "2026_Participant_Handbook.pdf", "2026_Kickoff_Deck.pdf"):
        copy(GS / pdf, f"{BRIEFS}/{pdf}")
    copy(GS / V2_ZIP, f"02_Source_Material/Data_v2_CURRENT/{V2_ZIP}")
    copy(GS / V1_ZIP, f"02_Source_Material/Data_v1_SUPERSEDED/{V1_ZIP}")
    note("02_Source_Material", "READ-ONLY. Restrict edit permissions on this folder after upload.\n"
         "Data_v2_CURRENT is the SME's 2026-09-29 archive and supersedes Data_v1_SUPERSEDED.\n"
         "Working copies of the CSVs are in 05_Build/code/data/raw.")


def build_analysis() -> None:
    copy(EDA / "report.html", "04_Analysis/UC4_Data_Profile.html")
    for f in (EDA / "tables").glob("*.csv"):
        copy(f, f"04_Analysis/Tables/{f.name}")
    for f in (EDA / "figures").glob("*.png"):
        copy(f, f"04_Analysis/Figures/{f.name}")
    for f in EDA.glob("*.py"):
        copy(f, f"04_Analysis/eda_code/{f.name}")
    shutil.copytree(EDA / "v1", OUT / "04_Analysis/v1_superseded", ignore=shutil.ignore_patterns("__pycache__"))
    copy(ROOT / "use-cases/uc4/DATA_ARCHITECTURE.md", "04_Analysis/UC4_Data_Architecture.md")
    convert_doc(ROOT / "use-cases/uc4/README.md", "04_Analysis/UC4_Guide.docx")
    convert_doc(ROOT / "use-cases/RECOMMENDATION.md", "04_Analysis/Team_Recommendation.docx")
    convert_doc(ROOT / "use-cases/VERIFICATION.md", "04_Analysis/Use_Case_Verification.docx")
    note("04_Analysis", "UC4_Data_Profile.html: open it from your synced OneDrive folder in a browser (onedrive.com downloads it rather than displaying it).\n"
         "UC4_Data_Architecture.md holds Mermaid diagrams; they render on GitHub/GitLab or at https://mermaid.live.\n"
         "eda_code regenerates the profile: uv run --no-project --with pandas --with matplotlib python uc4_eda.py\n"
         "(it expects the original workspace layout; see UC4_Guide.docx).")


def build_code() -> None:
    raw = OUT / "05_Build/code/data/raw"
    raw.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(GS / V2_ZIP) as z:
        for m in z.namelist():
            if m.lower().endswith(".csv"):
                (raw / Path(m).name).write_bytes(z.read(m))
    for sub, text in {
        "05_Build/code": "Temporary home for code until GitHub/GitLab is set up; move it there before the Friday submission.\n"
                         "One owner per file (OneDrive makes conflict copies on simultaneous edits).\n"
                         "Keep virtual environments (.venv), node_modules and __pycache__ OUT of this synced folder.",
        "05_Build/code/data/raw": "The seven v2 CSVs unzipped from 02_Source_Material/Data_v2_CURRENT. Never edit them.\n"
                                  "Five names carry a ' 1' download suffix; loaders should ignore it.",
        "05_Build/code/src": "Loader, SYNTH_V1 scoring engine (port 04_Analysis/eda_code/rules.py), question->answer, screen.",
        "05_Build/code/tests": "Start with the 72-of-72 verdict test (04_Analysis/eda_code/test_rules.py).",
        "05_Build/Screenshots": "Screenshots of working features for the deck and one-pager.",
    }.items():
        note(sub, text)


# ---------------------------------------------------------------- Word documents
START_HERE = """# UC4 R&D Data Source Unification: start here

**Demo:** Friday 2026-10-02, 7 minutes + 5 minutes Q&A. Submit the demo, deck, one-pager and repo link within **1 hour** afterwards ([handbook p. 3](02_Source_Material/Briefs_and_Rules/2026_Participant_Handbook.pdf)).

**What we build:** one place where a breeder asks a question about a trial, sees a PASS/HOLD/FAIL (green/amber/red) recommendation with every criterion, its value and threshold and the source rows, and records their own decision with a reason. Plain code scores; AI only interprets the question and phrases the answer from retrieved rows; the breeder decides.

## Team

| Person | Role | Hours / time zone |
|---|---|---|
| Harrison Fuller | _to assign_ | 6 AM–2 PM, UTC-8 |
| Carlos Diego Gomes | _to assign_ | |
| Jim Shilo | _to assign_ | |
| Matheus Carvalho | _to assign_ | |

Roles: **1 Data and rules**, **2 AI answers**, **3 Screen and log**, **4 Docs and demo**. Done-when checks and interface contracts: [Roles_and_Contracts.docx](01_Team/Roles_and_Contracts.docx).

## Where things are

| Folder | What goes there |
|---|---|
| 01_Team | Roles, decisions log, daily stand-up notes |
| 02_Source_Material | **Read-only.** Briefs, handbook, kickoff deck, v2 data (current) and v1 data (superseded) |
| 03_SME | Questions and answers with Diganta Adhikari (UC4 SME), saved emails |
| 04_Analysis | UC4 guide, data profile, inferred scoring rule, tables, figures, EDA code |
| 05_Build | Interface contracts, supported questions, code (until Git), screenshots |
| 06_AI_Prompt_Log | Prompt_Log.xlsx: log every AI prompt we keep, the same day |
| 07_Demo | Script, deck, one-pager, baseline timing, rehearsal recordings |
| 08_FINAL_Submission | Only what gets submitted; frozen after the demo |

## Rules for everyone

1. **Evidence first.** Every claim cites a CSV row, a brief/handbook page or a test result; otherwise label it *proposal* or *inferred*.
2. **AI does not decide.** Code sets the colour; the breeder makes the call.
3. **Missing evidence stays visible.** A missing value never passes a criterion.
4. **Log prompts the same day** in 06_AI_Prompt_Log.
5. **Edit in place.** Open Word/Excel files from this shared OneDrive folder (browser or synced desktop copy); don't download-edit-reupload. To sync it: open the share link, then **Add shortcut to My files**.

## This week

| Day | Goal |
|---|---|
| Tue 09-29 | Roles assigned, contracts agreed, SME follow-ups sent, data loaded |
| Wed 09-30 | Engine matches 72/72; first answers with citations; screen skeleton |
| Thu 10-01 | End-to-end flow incl. override log; baseline timing; deck draft; code moved to Git |
| Fri 10-02 | Rehearse within 7:00; demo; submit within 1 hour |
"""

STANDUP = """# Daily stand-up notes

15 minutes. Each person: done since last time, doing next, blocked by. Record decisions in 01_Team/Decisions_Log.xlsx.

""" + "\n".join(
    f"## {d}\n\n| Person | Done | Next | Blocked by |\n|---|---|---|---|\n"
    + "".join(f"| {p} | | | |\n" for p in ("Harrison Fuller", "Carlos Diego Gomes", "Jim Shilo", "Matheus Carvalho"))
    for d in ("Tue 2026-09-29", "Wed 2026-09-30", "Thu 2026-10-01", "Fri 2026-10-02")
)

SCRIPT = """# Demo script (7:00 hard limit)

Rehearse with a timer; record each run in 07_Demo/Rehearsals. Every claim must point to evidence (CSV row, page or test).

| Time | Section | Who | What is shown / said | Evidence |
|---|---|---|---|---|
| 0:00–0:45 | Problem | | The breeder pieces together trial, lab, genomics and operations data by hand | Brief p. 5 |
| 0:45–1:30 | Baseline | | Manual cross-file lookup for one trial: time and errors | Baseline_Timing.xlsx |
| 1:30–3:30 | Live demo: ask | | "Why is trial SYN-TR-0037 on hold?" → criteria table, resistant 30% < 50%, cited rows | trial_recommendations_synthetic.csv |
| 3:30–4:30 | Live demo: challenge and override | | Breeder records a decision with a reason; an override without a reason is rejected; log shows before/after | Override log |
| 4:30–5:30 | Trust | | Engine matches 72/72 supplied verdicts; thresholds labelled inferred; data contradictions shown as amber | Test output |
| 5:30–6:30 | How we worked with AI | | Prompt log across phases | Prompt_Log.xlsx |
| 6:30–7:00 | Close | | Value, limits, next steps | |

## Backup plan

If the live demo fails: screenshots in 05_Build/Screenshots and the most recent rehearsal recording.
"""

ONE_PAGER = """# UC4 R&D Data Source Unification: one-pager

**Team:** Harrison Fuller, Carlos Diego Gomes, Jim Shilo, Matheus Carvalho · **Repo:** _link_ · **Demo:** 2026-10-02

## Problem

_Two sentences, citing brief p. 5._

## What we built

_Three bullets: evidence view, cited answers, logged override._

## Evidence it works

| Check | Result |
|---|---|
| Engine vs supplied verdicts | _x / 72_ |
| Test questions answered correctly with citations | _x / y_ |
| Manual lookup vs app (time, correctness) | _from Baseline_Timing.xlsx_ |

## How AI was used across the lifecycle

_From the Prompt_Log.xlsx summary._

## Limits and next steps

_Inferred thresholds pending SME confirmation; synthetic data only; no live system connection._
"""

DECK = """---
title: UC4 R&D Data Source Unification
subtitle: HatchWorks AI x Syngenta hackathon, 2026-10-02
---

# The breeder's problem

- Placeholder: persona and fragmentation (brief p. 5)

# Baseline: doing it by hand

- Placeholder: time and errors from Baseline_Timing.xlsx

# What we built

- Placeholder: screenshot of the trial view

# Live demo

- Ask → evidence → challenge → override

# Why you can trust it

- Engine matches _x_/72 supplied verdicts
- Thresholds inferred from SYNTH_V1, pending SME confirmation
- Data contradictions shown, not hidden

# How we used AI across the lifecycle

- Placeholder: summary from Prompt_Log.xlsx

# Limits and next steps

- Placeholder
"""


def build_docs() -> None:
    pandoc(START_HERE, "00_START_HERE.docx")
    convert_doc(ROOT / "team/ROLES.md", "01_Team/Roles_and_Contracts.docx")
    pandoc(STANDUP, "01_Team/Standup_Notes.docx")
    convert_doc(ROOT / "team/SME_ANSWERS.md", "03_SME/SME_Questions_and_Answers.docx")
    note("03_SME/Emails", "Save each SME email here as .msg or .pdf, named YYYY-MM-DD_subject.")
    pandoc(SCRIPT, "07_Demo/Script_7min.docx")
    pandoc(ONE_PAGER, "07_Demo/One_Pager.docx")
    pandoc(DECK, "07_Demo/Deck.pptx", fmt_from="markdown")
    note("07_Demo/Rehearsals", "Recordings of each timed rehearsal, named YYYY-MM-DD_HHMM_runN_mm-ss.")
    note("08_FINAL_Submission", "Only the final demo recording, deck, one-pager and repo link. Freeze after submission.")


# ---------------------------------------------------------------- Excel workbooks
def build_decisions() -> None:
    wb = Workbook()
    ws = sheet(wb, "Decisions", ["Date", "Decision", "Why / evidence", "Decided by", "Status"], [
        ["2026-09-29", "Build UC4: R&D Data Source Unification", "Cleanest data, strongest AI fit, lowest sharing risk (Team_Recommendation.docx)", "Team", "Final"],
        ["2026-09-29", "v2 archive supersedes the kickoff (v1) archive", "SME email 2026-09-29 (03_SME)", "Team", "Pending SME confirmation"],
        ["2026-09-29", "Score with inferred rule SYNTH_V1; supplied verdicts are the baseline (must match 72/72)", "test_rules.py, rule_intervals.csv", "Team", "Thresholds pending SME confirmation"],
        ["2026-09-29", "Display PASS/HOLD/FAIL as green/amber/red", "Proposal; the file uses PASS/HOLD/FAIL only", "Team", "Proposal"],
        ["2026-09-29", "AI never sets the verdict; breeder makes the final call; overrides need a reason and are logged", "Brief p. 5; ROLES rule 2", "Team", "Final"],
        ["2026-09-29", "No live research-system connection; synthetic data only", "Brief p. 5 scope", "Team", "Final"],
        ["2026-09-29", "Single shared OneDrive folder until GitHub/GitLab is set up", "Team choice", "Team", "Final"],
    ], [12, 55, 55, 16, 26], first=True)
    dropdown(ws, "E", ["Proposal", "Final", "Pending SME confirmation", "Reversed"])
    wb.save(OUT / "01_Team/Decisions_Log.xlsx")


def build_contracts() -> None:
    cols = ["Field", "Type", "Required", "Meaning", "Example"]
    w = [26, 16, 10, 60, 30]
    wb = Workbook()
    sheet(wb, "Evidence row", cols, [
        ["source_file", "string", "yes", "CSV the value came from", "trial_recommendations_synthetic.csv"],
        ["row_id", "string", "yes", "Row identifier within the file (GUID or row number)", ""],
        ["trial_guid", "string | null", "no", "Null for lab and genomics rows (they have no trial key)", ""],
        ["material_guid", "string | null", "no", "Null for trial-level values", ""],
        ["field", "string", "yes", "Column name", "DISEASE_SCORE"],
        ["value", "string | number", "yes", "Original value, unchanged", "7.7"],
        ["uom", "string | null", "no", "Unit of measure if known", "t/ha"],
        ["date", "date | null", "no", "Date if the row has one", ""],
        ["flags[]", "list of string", "yes", "Consistency flags (e.g. operation outside trial year)", "[]"],
    ], w, first=True)
    sheet(wb, "Recommendation", cols, [
        ["trial_guid", "string", "yes", "Trial being scored", ""],
        ["verdict", "PASS | HOLD | FAIL", "yes", "Engine output (SYNTH_V1)", "HOLD"],
        ["colour", "green | amber | red", "yes", "Proposed display of verdict", "amber"],
        ["criteria[]", "list", "yes", "Each: field, value, threshold, met (true/false)", "RESISTANT_MATERIAL_PCT 30 >= 50 false"],
        ["knockout", "string | null", "no", "Knockout criterion that forced FAIL", "DISEASE_SCORE > 7"],
        ["reason", "string", "yes", "One-line reason naming every unmet criterion", ""],
        ["rule_version", "string", "yes", "Rule used", "SYNTH_V1 (inferred)"],
        ["supplied_verdict", "PASS | HOLD | FAIL", "yes", "TRIAL_RECOMMENDATION from the file; mismatch = bug", "HOLD"],
        ["evidence_row_ids[]", "list of string", "yes", "Evidence rows behind the verdict", ""],
        ["flags[]", "list of string", "yes", "Data contradictions to show as amber reasons", ""],
    ], w)
    sheet(wb, "Override record", cols, [
        ["trial_guid", "string", "yes", "Trial decided on", ""],
        ["material_guid", "string | null", "no", "Line, if the decision is per line", ""],
        ["recommendation", "Recommendation", "yes", "Copied at decision time; never modified", ""],
        ["decision", "string", "yes", "Breeder's final choice", "Advance"],
        ["reason", "string", "yes", "Required; an empty reason is rejected", ""],
        ["user", "string", "yes", "Breeder identity or demo alias", ""],
        ["timestamp", "ISO 8601 datetime", "yes", "When recorded (UTC)", "2026-10-01T15:04:00Z"],
    ], w)
    wb.save(OUT / "05_Build/Interface_Contracts.xlsx")


def build_questions() -> None:
    wb = Workbook()
    ws = sheet(wb, "Supported questions",
               ["#", "Question", "Type", "Expected answer (from data)", "Rows that must be cited", "Actual answer", "Checked by", "Result"], [
        [1, "Why is trial SYN-TR-0037 on hold?", "Explain verdict", "HOLD: resistant lines 30% < 50%; the supplied rationale omits this", "trial_recommendations_synthetic.csv SYN-TR-0037", "", "", ""],
        [2, "Why does trial SYN-TR-0003 pass?", "Explain verdict", "PASS: yield 10.79, moisture 16.8, disease 3.3, GBV mean 106.8, resistant 70%; all met", "trial_recommendations_synthetic.csv SYN-TR-0003", "", "", ""],
        [3, "Why did trial SYN-TR-0001 fail?", "Explain verdict", "FAIL: disease 7.7 > 7 knockout (yield 10.05 passes)", "trial_recommendations_synthetic.csv SYN-TR-0001", "", "", ""],
        [4, "Which trials fail on disease alone?", "List", "16 trials", "trial_recommendations_synthetic.csv", "", "", ""],
        [5, "Which trials fail on yield alone?", "List", "10 trials", "trial_recommendations_synthetic.csv", "", "", ""],
        [6, "Which trials read 'all criteria met' but are on hold?", "List", "SYN-TR-0037, 0038, 0046, 0052", "trial_recommendations_synthetic.csv", "", "", ""],
        [7, "Which lines are in trial SYN-TR-0037 and what are their genomic values?", "Lookup", "10 lines via observation links; show as context, not as verdict evidence", "observation, genomics", "", "", ""],
        [8, "Are there data problems with trial SYN-TR-00xx's operations?", "Consistency", "Flags from chronology checks (e.g. operation outside trial year)", "operations_synthetic 1.csv", "", "", ""],
        [9, "Tell me about SYN-TR-3", "Ambiguous ID", "Asks which trial is meant", "", "", "", ""],
        [10, "What is the yield of line SYN-MZ-00001?", "Missing evidence", "No line-level yield exists; says so", "", "", "", ""],
    ], [5, 45, 16, 50, 36, 40, 14, 10], first=True)
    dropdown(ws, "H", ["Pass", "Fail", "Not run"])
    wb.save(OUT / "05_Build/Supported_Questions.xlsx")


def build_prompt_log() -> None:
    wb = Workbook()
    ws = sheet(wb, "Log", ["Date", "Who", "Role", "Phase", "Tool / model", "Prompt (gist)", "Output / evidence", "Kept?", "Why (if edited/no)", "Cost"], [
        ["2026-09-29", "Harrison Fuller", "", "requirements", "Claude Code (Opus 5.5)",
         "We are choosing UC4. Create initial documentation to guide the group: split roles and start the shared prompt log.",
         "01_Team/Roles_and_Contracts.docx; this file", "yes", "", ""],
        ["2026-09-29", "Harrison Fuller", "Data and rules", "data", "Claude Code (Opus 5.5)",
         "We have new information [SME email + v2 zip]. Make evidence-based changes. Reverse-engineered PASS/HOLD/FAIL thresholds from 72 verdicts; rebuilt EDA for v2.",
         "03_SME/SME_Questions_and_Answers.docx; 04_Analysis/eda_code/rules.py; 04_Analysis/UC4_Data_Profile.html", "yes", "", ""],
        ["2026-09-29", "Harrison Fuller", "", "docs", "Claude Code (Opus 5.5)",
         "Provide the ideal shared folder structure and contents for UC4; build it locally (SharePoint, then OneDrive).",
         "This OneDrive folder", "yes", "", ""],
    ], [12, 18, 16, 14, 22, 60, 45, 9, 30, 10], first=True)
    dropdown(ws, "C", ["Data and rules", "AI answers", "Screen and log", "Docs and demo"])
    dropdown(ws, "D", ["requirements", "stories", "data", "prototype", "code", "test", "docs", "demo"])
    dropdown(ws, "H", ["yes", "edited", "no"])
    sheet(wb, "How to log", ["Rule"], [
        ["Why: judges score AI used across the lifecycle, not just for coding (handbook pp. 3-4). This log feeds the 'how we worked' slide."],
        ["Add a row the same day, newest at the bottom. Log prompts that produced something we kept or taught us something, including failures."],
        ["Phase: requirements, stories, data, prototype, code, test, docs, demo. Aim for entries in every phase by Thursday."],
        ["Prompt: the gist in one line; link the full text if long. Output / evidence: a file path, not a description."],
        ["Kept?: yes / edited / no, with why when edited or no. Human corrections are good evidence."],
        ["Cost: tokens or $ if the tool shows it (Portkey logs do); otherwise blank. Don't guess."],
        ["Never paste secrets, API keys or non-synthetic data."],
    ], [120])
    phases = ["requirements", "stories", "data", "prototype", "code", "test", "docs", "demo"]
    ws = sheet(wb, "Roll-up", ["Phase", "# entries", "Best example", "Lesson"],
               [[p, f'=COUNTIF(Log!D:D,"{p}")', "", ""] for p in phases], [16, 10, 50, 50])
    (OUT / "06_AI_Prompt_Log").mkdir(exist_ok=True)
    wb.save(OUT / "06_AI_Prompt_Log/Prompt_Log.xlsx")


def build_baseline() -> None:
    wb = Workbook()
    ws = sheet(wb, "Timing", ["Task", "Trial / line", "Method", "Person", "Minutes", "Answer given", "Correct?", "Notes"], [
        [t, tid, m, "", "", "", "", ""]
        for t, tid in (("Explain the verdict and every criterion vs threshold", "SYN-TR-0037"),
                       ("Explain the verdict and every criterion vs threshold", "SYN-TR-0001"),
                       ("List the lines in the trial with their genomic values", "SYN-TR-0003"),
                       ("Find data contradictions in the trial's operations", ""))
        for m in ("Manual (CSV files)", "App")
    ], [48, 14, 20, 16, 10, 40, 10, 40], first=True)
    dropdown(ws, "C", ["Manual (CSV files)", "App"])
    dropdown(ws, "G", ["yes", "partly", "no"])
    sheet(wb, "Method", ["Rule"], [
        ["Same person must not do manual and app runs for the same task back to back (memory bias); alternate or swap people."],
        ["Manual: only the raw CSVs in 05_Build/code/data/raw, any spreadsheet tool. App: the demo build."],
        ["Stop the timer when the person states the full answer. Check correctness against Supported_Questions.xlsx."],
        ["Do not claim a shorter breeding cycle; report lookup time and correctness only (UC4 guide)."],
    ], [120])
    wb.save(OUT / "07_Demo/Baseline_Timing.xlsx")


def main() -> None:
    if OUT.exists():
        if not FORCE:
            sys.exit(f"{OUT} exists; it may hold teammates' edits. Re-run with --force to replace it.")
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    build_source()
    build_analysis()
    build_code()
    build_docs()
    build_decisions()
    build_contracts()
    build_questions()
    build_prompt_log()
    build_baseline()
    print(OUT)


if __name__ == "__main__":
    main()
